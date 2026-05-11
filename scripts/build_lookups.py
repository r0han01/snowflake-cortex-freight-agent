"""
Build the 6 lookup CSVs from the official FAF5 metadata workbook.

Source : ../../FAF5.7.1_State_2018-2024/FAF5_metadata.xlsx
Output : ../lookups/lookup_{states,sctg,modes,trade_types,dist_bands,foreign_regions}.csv

We extract directly from the official codebook so the codes line up byte-for-byte with the
fact-table values. Re-runnable; overwrites in place.
"""

import csv
import os
import sys
import openpyxl

# ---- Paths -------------------------------------------------------------------
HERE          = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT  = os.path.dirname(HERE)
METADATA_XLSX = os.path.normpath(os.path.join(
    PROJECT_ROOT, "..", "FAF5.7.1_State_2018-2024", "FAF5_metadata.xlsx"
))
LOOKUPS_DIR   = os.path.join(PROJECT_ROOT, "lookups")
os.makedirs(LOOKUPS_DIR, exist_ok=True)

# ---- Schemas -----------------------------------------------------------------
# (sheet_name, output_filename, header_cols, expected_min_rows, header_skip_rows)
SHEETS = [
    ("State",              "lookup_states.csv",          ["state_code", "state_name"],          51, 1),
    ("Commodity (SCTG2)",  "lookup_sctg.csv",            ["sctg_code", "commodity_name"],       42, 1),
    ("Mode",               "lookup_modes.csv",           ["mode_code", "mode_name"],             8, 1),
    ("Trade Type",         "lookup_trade_types.csv",     ["trade_code", "trade_type_name"],      3, 1),
    ("Distance Band",      "lookup_dist_bands.csv",      ["band_code", "band_label"],            8, 1),
    ("FAF Zone (Foreign)", "lookup_foreign_regions.csv", ["region_code", "region_name"],         8, 1),
]


def clean(v):
    """Normalize a cell: strip, remove Excel CR artifacts, return '' if None."""
    if v is None:
        return ""
    s = str(v).strip()
    # Excel sometimes embeds literal _x000D_ for carriage returns
    s = s.replace("_x000D_", "").replace("\r", "")
    return s


# Sheet-specific name overrides. The metadata workbook ships verbose descriptions for some sheets
# (e.g. "Domestic flows (freight shipments moved from...)") — we prefer short canonical labels in
# the agent's surface area. Codes are unchanged; only the human-readable name is rewritten.
NAME_OVERRIDES = {
    "Trade Type": {
        "1": "Domestic",
        "2": "Import",
        "3": "Export",
    },
}


def extract_sheet(wb, sheet_name, columns, expected_min, skip_header_rows):
    if sheet_name not in wb.sheetnames:
        raise SystemExit(f"❌ Sheet '{sheet_name}' not found in metadata workbook")
    ws = wb[sheet_name]
    overrides = NAME_OVERRIDES.get(sheet_name, {})
    rows_out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < skip_header_rows:
            continue
        if not any(c is not None and clean(c) != "" for c in row):
            continue  # skip fully empty rows
        code = clean(row[0])
        name = clean(row[1]) if len(row) > 1 else ""
        if not code or not name:
            continue
        # Apply override if defined for this sheet+code
        if code in overrides:
            name = overrides[code]
        rows_out.append((code, name))
    return rows_out


def write_csv(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def main():
    if not os.path.isfile(METADATA_XLSX):
        sys.exit(f"❌ Metadata workbook not found: {METADATA_XLSX}")
    print(f"Reading metadata: {METADATA_XLSX}")
    wb = openpyxl.load_workbook(METADATA_XLSX, data_only=True)
    print(f"Output dir     : {LOOKUPS_DIR}")
    print()
    print(f"{'File':<35} {'Rows':>5}   {'Status':<10}  Sample")
    print("-" * 100)

    all_ok = True
    summary = []
    for sheet_name, filename, columns, expected, skip in SHEETS:
        rows = extract_sheet(wb, sheet_name, columns, expected, skip)
        path = os.path.join(LOOKUPS_DIR, filename)
        write_csv(path, columns, rows)
        status = "OK" if len(rows) >= expected else f"⚠️ <{expected}"
        sample = ", ".join([f"({c}={n})" for c, n in rows[:2]])
        if len(rows) < expected:
            all_ok = False
        summary.append((filename, len(rows), status, sample))
        print(f"{filename:<35} {len(rows):>5}   {status:<10}  {sample}")

    print()
    print("-" * 100)
    total = sum(s[1] for s in summary)
    print(f"Total rows across all 6 lookups: {total}")
    print()
    if all_ok:
        print("✅ All lookups built successfully and meet expected minimum row counts.")
    else:
        print("⚠️ Some lookups have fewer rows than expected — review above.")
        sys.exit(1)


if __name__ == "__main__":
    main()
